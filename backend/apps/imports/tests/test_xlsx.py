"""XLSX worksheet parsing, numeric codes and formula rejection."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from apps.imports.services import commit_batch, preview_batch
from apps.imports.tests.conftest import IMPORT_FILES, PREVIEW_NOW, TEMPLATES_DIR, make_batch, read_csv
from apps.scheduling.models import ClassSession

pytestmark = pytest.mark.django_db


def build_workbook(path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for file_name in IMPORT_FILES:
        rows = read_csv(TEMPLATES_DIR / file_name)
        worksheet = workbook.create_sheet(Path(file_name).stem)
        for row in rows:
            worksheet.append(row)
    workbook.save(path)
    return path


def _find_issue(result, *, sheet: str, row: int, field: str, code: str):
    for issue in result.issues:
        if (issue.sheet, issue.row, issue.field, issue.code) == (sheet, row, field, code):
            return issue
    return None


class TestXlsx:
    def test_twelve_sheets_have_the_same_result_as_csv(self, school, academic_staff, tmp_path) -> None:
        workbook_path = build_workbook(tmp_path / "templates.xlsx")
        batch = make_batch(school, files=[workbook_path])

        preview = preview_batch(batch, now=PREVIEW_NOW)
        assert not preview.has_errors, preview.errors
        assert preview.expected_new_sessions == 40

        commit = commit_batch(batch, academic_staff)

        assert commit.committed
        assert ClassSession.objects.count() == 40

    def test_numeric_code_cell_requires_text(self, school, tmp_path) -> None:
        workbook_path = tmp_path / "numeric.xlsx"
        workbook = Workbook()
        workbook.remove(workbook.active)
        for file_name in IMPORT_FILES:
            rows = read_csv(TEMPLATES_DIR / file_name)
            worksheet = workbook.create_sheet(Path(file_name).stem)
            for index, row in enumerate(rows, start=1):
                worksheet.append(row)
                if file_name == "teachers.csv" and index == 2:
                    worksheet.cell(row=2, column=1, value=1)
        workbook.save(workbook_path)
        batch = make_batch(school, files=[workbook_path])

        result = preview_batch(batch, now=PREVIEW_NOW)

        issue = _find_issue(result, sheet="teachers", row=2, field="teacher_code", code="CODE_NOT_TEXT")
        assert issue is not None
        assert issue.file == "numeric.xlsx"

    def test_formula_cell_is_rejected(self, school, tmp_path) -> None:
        workbook_path = build_workbook(tmp_path / "formula.xlsx")
        workbook = Workbook()
        workbook.remove(workbook.active)
        for file_name in IMPORT_FILES:
            rows = read_csv(TEMPLATES_DIR / file_name)
            worksheet = workbook.create_sheet(Path(file_name).stem)
            for row in rows:
                worksheet.append(row)
            if file_name == "terms.csv":
                worksheet["B2"] = '="虚构演示学校2026秋季学期"'
        workbook.save(workbook_path)
        batch = make_batch(school, files=[workbook_path])

        result = preview_batch(batch, now=PREVIEW_NOW)

        issue = _find_issue(result, sheet="terms", row=2, field="term_name", code="FORMULA_NOT_ALLOWED")
        assert issue is not None
        assert issue.file == "formula.xlsx"
