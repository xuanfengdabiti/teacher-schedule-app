"""Fixed-template CSV/XLSX parsing.

The parser deliberately does not infer columns or repair spreadsheet values.
It produces typed rows plus stable, row-addressable issues; cross-row business
rules live in :mod:`apps.imports.validation`.
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from openpyxl import load_workbook

FORMULA_NOT_ALLOWED = "FORMULA_NOT_ALLOWED"
CODE_NOT_TEXT = "CODE_NOT_TEXT"
PLACEHOLDER_NOT_ALLOWED = "PLACEHOLDER_NOT_ALLOWED"

_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_TIME_RE = re.compile(r"^\d{2}:\d{2}$")
_INTEGER_RE = re.compile(r"^-?\d+$")
_HOURS_RE = re.compile(r"^\d+(?:\.\d)?$")
_DATETIME_OFFSET_RE = re.compile(r"(?:Z|[+-]\d{2}:\d{2})$")
_PLACEHOLDERS = {"null", "-", "无"}


@dataclass(frozen=True)
class FieldSpec:
    name: str
    kind: str
    required: bool = True
    max_length: int | None = None


@dataclass(frozen=True)
class FileSpec:
    logical_name: str
    identity_fields: tuple[str, ...]
    fields: tuple[FieldSpec, ...]

    @property
    def field_names(self) -> tuple[str, ...]:
        return tuple(field.name for field in self.fields)

    @property
    def field_specs(self) -> dict[str, FieldSpec]:
        return {field.name: field for field in self.fields}


def _f(name: str, kind: str, required: bool = True, max_length: int | None = None) -> FieldSpec:
    return FieldSpec(name=name, kind=kind, required=required, max_length=max_length)


FILE_SPECS: dict[str, FileSpec] = {
    "terms": FileSpec(
        "terms",
        ("term_code",),
        (
            _f("term_code", "code", max_length=64),
            _f("term_name", "text", max_length=200),
            _f("start_date", "date"),
            _f("end_date", "date"),
            _f("first_week_monday", "date"),
            _f("teaching_week_count", "integer"),
            _f("timezone", "text", max_length=64),
            _f("hour_minutes", "integer"),
        ),
    ),
    "periods": FileSpec(
        "periods",
        ("term_code", "period_number"),
        (
            _f("term_code", "code", max_length=64),
            _f("period_number", "integer"),
            _f("start_time", "time"),
            _f("end_time", "time"),
            _f("block_code", "code", max_length=64),
        ),
    ),
    "teachers": FileSpec(
        "teachers",
        ("teacher_code",),
        (
            _f("teacher_code", "code", max_length=64),
            _f("formal_name", "text", max_length=100),
            _f("department", "text", False, 100),
            _f("email", "email", False, 254),
            _f("phone", "text", False, 32),
            _f("phone_visibility", "text", max_length=32),
            _f("wechat", "text", False, 64),
            _f("wechat_visibility", "text", max_length=32),
            _f("is_active", "bool"),
        ),
    ),
    "classrooms": FileSpec(
        "classrooms",
        ("classroom_code",),
        (
            _f("classroom_code", "code", max_length=64),
            _f("campus", "text", max_length=100),
            _f("building", "text", max_length=100),
            _f("room_no", "text", max_length=64),
            _f("capacity", "integer"),
            _f("equipment", "codes", False, 64),
            _f("is_active", "bool"),
        ),
    ),
    "offerings": FileSpec(
        "offerings",
        ("term_code", "offering_code"),
        (
            _f("term_code", "code", max_length=64),
            _f("offering_code", "code", max_length=64),
            _f("course_code", "code", max_length=64),
            _f("course_name", "text", max_length=200),
            _f("class_code", "code", max_length=64),
            _f("class_name", "text", max_length=200),
            _f("target_hours", "hours"),
            _f("lead_teacher_code", "code", max_length=64),
            _f("teacher_codes", "codes", max_length=64),
            _f("required_equipment", "codes", False, 64),
        ),
    ),
    "enrollments": FileSpec(
        "enrollments",
        ("enrollment_code",),
        (
            _f("enrollment_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("offering_code", "code", max_length=64),
            _f("student_code", "code", max_length=64),
            _f("effective_from", "date"),
            _f("effective_until", "date", False),
        ),
    ),
    "calendar_exceptions": FileSpec(
        "calendar_exceptions",
        ("exception_code",),
        (
            _f("exception_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("date", "date"),
            _f("kind", "text", max_length=16),
            _f("source_teaching_date", "date", False),
            _f("reason", "text", max_length=2000),
        ),
    ),
    "rules": FileSpec(
        "rules",
        ("term_code", "rule_code"),
        (
            _f("rule_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("offering_code", "code", max_length=64),
            _f("weekday", "integer"),
            _f("start_period", "integer"),
            _f("end_period", "integer"),
            _f("weeks", "integer_list"),
            _f("teacher_codes", "codes", max_length=64),
            _f("classroom_code", "code", max_length=64),
            _f("planned_hours", "hours"),
        ),
    ),
    "sessions": FileSpec(
        "sessions",
        ("external_session_id",),
        (
            _f("external_session_id", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("offering_code", "code", max_length=64),
            _f("origin_rule_code", "code", False, 64),
            _f("origin_teaching_date", "date", False),
            _f("origin_start_period", "integer", False),
            _f("session_date", "date"),
            _f("start_period", "integer"),
            _f("end_period", "integer"),
            _f("start_time", "time", False),
            _f("end_time", "time", False),
            _f("classroom_code", "code", max_length=64),
            _f("teacher_codes", "codes", max_length=64),
            _f("status", "text", max_length=16),
            _f("planned_hours", "hours"),
            _f("actual_hours", "hours", False),
            _f("completed_by_teacher_code", "code", False, 64),
            _f("completed_at", "datetime", False),
            _f("completion_ref", "text", False, 200),
            _f("change_ref", "text", False, 200),
        ),
    ),
    "room_blocks": FileSpec(
        "room_blocks",
        ("block_code",),
        (
            _f("block_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("classroom_code", "code", max_length=64),
            _f("start_at", "datetime"),
            _f("end_at", "datetime"),
            _f("reason", "text", max_length=2000),
        ),
    ),
    "resource_unavailability": FileSpec(
        "resource_unavailability",
        ("block_code",),
        (
            _f("block_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("resource_type", "text", max_length=16),
            _f("resource_code", "code", max_length=64),
            _f("start_at", "datetime"),
            _f("end_at", "datetime"),
            _f("reason", "text", max_length=2000),
        ),
    ),
    "coverage": FileSpec(
        "coverage",
        ("coverage_code",),
        (
            _f("coverage_code", "code", max_length=64),
            _f("term_code", "code", max_length=64),
            _f("kind", "text", max_length=32),
            _f("scope_key", "text", max_length=200),
            _f("from_date", "date"),
            _f("until_date", "date"),
            _f("completeness", "text", max_length=16),
            _f("source_updated_at", "datetime"),
            _f("confirmation_ref", "text", False, 200),
        ),
    ),
}


@dataclass(frozen=True)
class ImportIssue:
    file: str
    sheet: str
    row: int
    field: str
    value: str
    code: str
    message: str
    level: str = "error"

    def as_dict(self) -> dict[str, Any]:
        return {
            "file": self.file,
            "sheet": self.sheet,
            "row": self.row,
            "field": self.field,
            "value": self.value,
            "code": self.code,
            "message": self.message,
            "level": self.level,
        }


@dataclass
class ParsedRow:
    file: str
    sheet: str
    row_number: int
    logical_name: str
    values: dict[str, Any]

    @property
    def source_key(self) -> str:
        return f"{self.file}\x1f{self.sheet}\x1f{self.row_number}"

    @property
    def display_location(self) -> str:
        if self.sheet:
            return f"{self.file}[{self.sheet}]:{self.row_number}"
        return f"{self.file}:{self.row_number}"


@dataclass
class ParsedFile:
    file_name: str
    sheet: str
    logical_name: str
    rows: list[ParsedRow] = field(default_factory=list)

    @property
    def display_key(self) -> str:
        return f"{self.file_name} / {self.sheet}" if self.sheet else self.file_name


@dataclass
class ParseResult:
    files: list[ParsedFile]
    issues: list[ImportIssue]

    @property
    def rows(self) -> list[ParsedRow]:
        return [row for parsed_file in self.files for row in parsed_file.rows]

    def rows_for(self, logical_name: str) -> list[ParsedRow]:
        return [row for row in self.rows if row.logical_name == logical_name]


def truncate_value(value: Any) -> str:
    """Return the original value for an error report, capped at 200 characters."""
    if value is None:
        return ""
    if isinstance(value, bool):
        rendered = "true" if value else "false"
    elif isinstance(value, (date, time, datetime)):
        rendered = value.isoformat()
    elif isinstance(value, (list, tuple)):
        rendered = "|".join(str(item) for item in value)
    else:
        rendered = str(value)
    return rendered[:200]


def make_issue(
    *,
    file: str,
    sheet: str,
    row: int,
    field: str,
    value: Any,
    code: str,
    message: str,
    level: str = "error",
) -> ImportIssue:
    return ImportIssue(
        file=file,
        sheet=sheet,
        row=row,
        field=field,
        value=truncate_value(value),
        code=code,
        message=message,
        level=level,
    )


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_code(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if not isinstance(value, str):
        return None, _field_issue(
            field_spec,
            value,
            CODE_NOT_TEXT,
            "编号列必须使用文本格式，以保留前导零；请在 Excel 中将该列设为文本后重新导出。",
        )
    value = value.strip()
    issue = _length_issue(field_spec, value)
    return value, issue


def _parse_text(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if not isinstance(value, str):
        return None, _field_issue(field_spec, value, "INVALID_TYPE", "该字段必须是文本；请检查单元格格式。")
    value = value.strip()
    issue = _length_issue(field_spec, value)
    return value, issue


def _parse_email(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    text, issue = _parse_text(field_spec, value)
    if issue or not text:
        return text, issue
    try:
        validate_email(text)
    except ValidationError:
        return None, _field_issue(field_spec, value, "INVALID_EMAIL", "邮箱格式无效；请填写可投递的邮箱或留空。")
    return text, None


def _parse_date(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if not isinstance(value, str):
        return None, _field_issue(field_spec, value, "INVALID_DATE", "日期必须是 YYYY-MM-DD 文本。")
    text = value.strip()
    if not _DATE_RE.fullmatch(text):
        return None, _field_issue(field_spec, value, "INVALID_DATE", "日期格式应为 YYYY-MM-DD；请修正后重试。")
    try:
        return date.fromisoformat(text), None
    except ValueError:
        return None, _field_issue(field_spec, value, "INVALID_DATE", "日期不存在；请核对月份和日期。")


def _parse_time(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0), None
    if not isinstance(value, str):
        return None, _field_issue(field_spec, value, "INVALID_TIME", "时间必须是 HH:mm 文本。")
    text = value.strip()
    if not _TIME_RE.fullmatch(text):
        return None, _field_issue(field_spec, value, "INVALID_TIME", "时间格式应为 HH:mm（24 小时制）。")
    try:
        return time.fromisoformat(text), None
    except ValueError:
        return None, _field_issue(field_spec, value, "INVALID_TIME", "时间不存在；请核对小时和分钟。")


def _parse_datetime(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if not isinstance(value, str):
        return None, _field_issue(
            field_spec,
            value,
            "TZ_REQUIRED",
            "时间点必须是以 ISO 8601 文本给出并带时区偏移，例如 2026-10-26T08:00:00+08:00。",
        )
    text = value.strip()
    if not _DATETIME_OFFSET_RE.search(text):
        return None, _field_issue(
            field_spec, value, "TZ_REQUIRED", "datetime 必须包含具体时区偏移，不能使用无时区的单元格。"
        )
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return None, _field_issue(field_spec, value, "INVALID_DATETIME", "datetime 不是有效的 ISO 8601 值。")
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None, _field_issue(field_spec, value, "TZ_REQUIRED", "datetime 必须包含具体时区偏移。")
    return parsed, None


def _parse_integer(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if isinstance(value, bool):
        return None, _field_issue(field_spec, value, "INVALID_INTEGER", "该字段必须是十进制整数。")
    if isinstance(value, int):
        return value, None
    if not isinstance(value, str):
        return None, _field_issue(
            field_spec, value, "INVALID_INTEGER", "该字段必须是十进制整数，不能使用小数或科学计数法。"
        )
    text = value.strip()
    if not _INTEGER_RE.fullmatch(text):
        return None, _field_issue(
            field_spec, value, "INVALID_INTEGER", "该字段必须是十进制整数，不能使用小数或科学计数法。"
        )
    return int(text), None


def _parse_hours(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if isinstance(value, bool):
        return None, _field_issue(field_spec, value, "INVALID_HOURS", "学时必须是最多 1 位小数的十进制非负数。")
    if isinstance(value, (int, float, Decimal)):
        text = str(value)
    elif isinstance(value, str):
        text = value.strip()
    else:
        return None, _field_issue(field_spec, value, "INVALID_HOURS", "学时必须是最多 1 位小数的十进制非负数。")
    if not _HOURS_RE.fullmatch(text):
        return None, _field_issue(
            field_spec, value, "INVALID_HOURS", "学时必须是最多 1 位小数的十进制非负数，例如 2.0。"
        )
    try:
        parsed = Decimal(text)
    except InvalidOperation:
        return None, _field_issue(field_spec, value, "INVALID_HOURS", "学时格式无效。")
    if parsed < 0:
        return None, _field_issue(field_spec, value, "INVALID_HOURS", "学时不能为负数。")
    if parsed > Decimal("99999.9"):
        return None, _field_issue(field_spec, value, "INVALID_HOURS", "学时超出 Decimal(6,1) 可存储范围。")
    return parsed, None


def _parse_bool(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if isinstance(value, bool):
        return value, None
    if isinstance(value, str):
        text = value.strip()
        if text == "true":
            return True, None
        if text == "false":
            return False, None
    return None, _field_issue(field_spec, value, "INVALID_BOOL", "布尔值只能填写 true 或 false，不能使用是/否、0/1。")


def _parse_codes(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if not isinstance(value, str):
        return [], _field_issue(
            field_spec, value, "INVALID_CODES", "代码集合必须是以 | 分隔的文本，不能使用数字单元格。"
        )
    parts = [part.strip() for part in value.split("|")]
    if any(not part for part in parts):
        return [], _field_issue(field_spec, value, "INVALID_CODES", "代码集合中存在空项；请删除多余的分隔符。")
    too_long = [part for part in parts if field_spec.max_length and len(part) > field_spec.max_length]
    if too_long:
        return [], _field_issue(
            field_spec, value, "CODE_TOO_LONG", f"代码长度不得超过 {field_spec.max_length} 个字符。"
        )
    return sorted(set(parts)), None


def _parse_integer_list(field_spec: FieldSpec, value: Any) -> tuple[Any, ImportIssue | None]:
    if not isinstance(value, str):
        return [], _field_issue(
            field_spec, value, "INVALID_INTEGER_LIST", "教学周必须是 | 分隔的整数集合，例如 1|3|5。"
        )
    parts = [part.strip() for part in value.split("|")]
    if any(not _INTEGER_RE.fullmatch(part) for part in parts):
        return [], _field_issue(field_spec, value, "INVALID_INTEGER_LIST", "教学周必须是 | 分隔的十进制整数集合。")
    return sorted({int(part) for part in parts}), None


def _length_issue(field_spec: FieldSpec, value: str) -> ImportIssue | None:
    if field_spec.max_length and len(value) > field_spec.max_length:
        return _field_issue(
            field_spec,
            value,
            "TOO_LONG",
            f"字段长度 {len(value)} 超过上限 {field_spec.max_length} 个字符。",
        )
    return None


def _field_issue(field_spec: FieldSpec, value: Any, code: str, message: str) -> ImportIssue:
    return make_issue(
        file="",
        sheet="",
        row=0,
        field=field_spec.name,
        value=value,
        code=code,
        message=message,
    )


def _parse_value(field_spec: FieldSpec, value: Any) -> tuple[Any, tuple[str, str] | None]:
    if _is_blank(value):
        if field_spec.required:
            return None, ("REQUIRED", "必填字段为空；请填写规范值，不允许用 null、- 或“无”代替空值。")
        return None, None

    if isinstance(value, str) and value.strip().casefold() in _PLACEHOLDERS:
        return None, (
            PLACEHOLDER_NOT_ALLOWED,
            "可空字段请直接留空，不要填写 null、- 或“无”；这些不是合法值。",
        )

    parser = {
        "code": _parse_code,
        "text": _parse_text,
        "email": _parse_email,
        "date": _parse_date,
        "time": _parse_time,
        "datetime": _parse_datetime,
        "integer": _parse_integer,
        "hours": _parse_hours,
        "bool": _parse_bool,
        "codes": _parse_codes,
        "integer_list": _parse_integer_list,
    }.get(field_spec.kind)
    if parser is None:
        raise ValueError(f"未知字段类型：{field_spec.kind}")
    parsed, issue = parser(field_spec, value)
    if issue is not None:
        return None, (issue.code, issue.message)
    return parsed, None


def _normalized_cell_array(values: list[Any]) -> tuple[Any, ...]:
    while values and values[-1] in (None, ""):
        values.pop()
    return tuple(values)


def _parse_table(
    *,
    file_name: str,
    sheet: str,
    logical_name: str,
    header_values: list[Any],
    data_rows: list[tuple[int, list[Any], list[bool]]],
) -> tuple[ParsedFile, list[ImportIssue]]:
    spec = FILE_SPECS[logical_name]
    issues: list[ImportIssue] = []
    parsed_file = ParsedFile(file_name=file_name, sheet=sheet, logical_name=logical_name)
    headers = _normalized_cell_array(list(header_values))
    expected = spec.field_names

    if not headers or all(value in (None, "") for value in headers):
        issues.append(
            make_issue(
                file=file_name,
                sheet=sheet,
                row=1,
                field="",
                value="",
                code="EMPTY_FILE",
                message="文件或工作表没有表头；请使用固定模板并保留第 1 行字段名。",
            )
        )
        return parsed_file, issues

    header_strings = [str(value) if value is not None else "" for value in headers]
    missing = [name for name in expected if name not in header_strings]
    extra = [name for name in header_strings if name not in expected]
    duplicates = sorted({name for name in header_strings if header_strings.count(name) > 1})
    header_ok = not missing and not extra and not duplicates
    if not header_ok:
        for name in missing:
            issues.append(
                make_issue(
                    file=file_name,
                    sheet=sheet,
                    row=1,
                    field=name,
                    value="",
                    code="HEADER_FIELDS_MISMATCH",
                    message=f"缺少规范字段 {name}；请使用固定模板，不要删除或改名。",
                )
            )
        for name in extra:
            issues.append(
                make_issue(
                    file=file_name,
                    sheet=sheet,
                    row=1,
                    field=name,
                    value=name,
                    code="HEADER_FIELDS_MISMATCH",
                    message=f"发现规范之外的字段 {name}；请删除多余列并核对拼写。",
                )
            )
        for name in duplicates:
            issues.append(
                make_issue(
                    file=file_name,
                    sheet=sheet,
                    row=1,
                    field=name,
                    value=name,
                    code="HEADER_FIELDS_MISMATCH",
                    message=f"表头字段 {name} 重复；每个规范字段只能出现一次。",
                )
            )
        return parsed_file, issues

    header_index = {name: index for index, name in enumerate(header_strings)}
    for row_number, row_values, formula_flags in data_rows:
        cells = list(row_values)
        while (
            len(cells) > len(headers)
            and cells[-1] in (None, "")
            and not (len(formula_flags) >= len(cells) and formula_flags[len(cells) - 1])
        ):
            cells.pop()
        if all(value in (None, "") for value in cells):
            continue
        if len(cells) != len(headers):
            issues.append(
                make_issue(
                    file=file_name,
                    sheet=sheet,
                    row=row_number,
                    field="",
                    value=",".join(truncate_value(value) for value in cells),
                    code="COLUMN_COUNT_MISMATCH",
                    message=f"该行有 {len(cells)} 列，但表头有 {len(headers)} 列；请补齐或删除多余单元格。",
                )
            )

        values: dict[str, Any] = {}
        row_issues: list[tuple[str, str, Any]] = []
        for field_spec in spec.fields:
            index = header_index[field_spec.name]
            raw_value = cells[index] if index < len(cells) else None
            if index < len(formula_flags) and formula_flags[index]:
                row_issues.append(
                    (
                        field_spec.name,
                        FORMULA_NOT_ALLOWED,
                        "不允许使用公式单元格；请粘贴为静态值后重新导出，避免依赖 Excel 缓存计算结果。",
                    )
                )
                values[field_spec.name] = None
                continue
            parsed, error = _parse_value(field_spec, raw_value)
            values[field_spec.name] = parsed
            if error is not None:
                row_issues.append((field_spec.name, error[0], error[1]))

        parsed_row = ParsedRow(
            file=file_name,
            sheet=sheet,
            row_number=row_number,
            logical_name=logical_name,
            values=values,
        )
        parsed_file.rows.append(parsed_row)
        for field_name, code, message in row_issues:
            index = header_index.get(field_name)
            raw = cells[index] if index is not None and index < len(cells) else None
            issues.append(
                make_issue(
                    file=file_name,
                    sheet=sheet,
                    row=row_number,
                    field=field_name,
                    value=raw,
                    code=code,
                    message=message,
                )
            )
    return parsed_file, issues


def _read_csv(path: Path, file_name: str) -> tuple[list[ParsedFile], list[ImportIssue]]:
    logical_name = path.stem
    if logical_name not in FILE_SPECS:
        return [], [
            make_issue(
                file=file_name,
                sheet="",
                row=0,
                field="",
                value=file_name,
                code="UNKNOWN_FILE",
                message=f"文件名 {path.name} 不在 12 个固定模板中；请使用规范文件名。",
            )
        ]
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            try:
                header = next(reader)
            except StopIteration:
                return [], [
                    make_issue(
                        file=file_name,
                        sheet="",
                        row=1,
                        field="",
                        value="",
                        code="EMPTY_FILE",
                        message="CSV 文件为空；请保留表头和数据行。",
                    )
                ]
            header_line = reader.line_num
            if header_line != 1:
                header_line = 1
            previous_line = header_line
            data_rows: list[tuple[int, list[Any], list[bool]]] = []
            for row in reader:
                row_number = previous_line + 1
                previous_line = reader.line_num
                data_rows.append((row_number, list(row), [False] * len(row)))
    except UnicodeDecodeError:
        return [], [
            make_issue(
                file=file_name,
                sheet="",
                row=0,
                field="",
                value=file_name,
                code="INVALID_ENCODING",
                message="CSV 必须以 UTF-8 编码（可带 BOM）；请重新导出为 UTF-8。",
            )
        ]
    except OSError as exc:
        return [], [
            make_issue(
                file=file_name,
                sheet="",
                row=0,
                field="",
                value=file_name,
                code="FILE_READ_FAILED",
                message=f"无法读取文件：{exc}",
            )
        ]
    parsed_file, issues = _parse_table(
        file_name=file_name,
        sheet="",
        logical_name=logical_name,
        header_values=list(header),
        data_rows=data_rows,
    )
    return [parsed_file], issues


def _read_xlsx(path: Path, file_name: str) -> tuple[list[ParsedFile], list[ImportIssue]]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises several format-specific exceptions
        return [], [
            make_issue(
                file=file_name,
                sheet="",
                row=0,
                field="",
                value=file_name,
                code="INVALID_XLSX",
                message=f"无法读取 XLSX：{exc}",
            )
        ]

    parsed_files: list[ParsedFile] = []
    issues: list[ImportIssue] = []
    try:
        for worksheet in workbook.worksheets:
            logical_name = worksheet.title.strip()
            if logical_name not in FILE_SPECS:
                continue
            rows_iterator = worksheet.iter_rows()
            try:
                header_cells = next(rows_iterator)
            except StopIteration:
                issues.append(
                    make_issue(
                        file=file_name,
                        sheet=worksheet.title,
                        row=1,
                        field="",
                        value="",
                        code="EMPTY_FILE",
                        message=f"工作表 {worksheet.title} 为空；请保留表头和数据行。",
                    )
                )
                continue
            header_values = [cell.value for cell in header_cells]
            data_rows: list[tuple[int, list[Any], list[bool]]] = []
            for row_number, cells in enumerate(rows_iterator, start=2):
                values = [cell.value for cell in cells]
                formula_flags = [cell.data_type == "f" for cell in cells]
                data_rows.append((row_number, values, formula_flags))
            parsed_file, table_issues = _parse_table(
                file_name=file_name,
                sheet=worksheet.title,
                logical_name=logical_name,
                header_values=header_values,
                data_rows=data_rows,
            )
            parsed_files.append(parsed_file)
            issues.extend(table_issues)
    finally:
        workbook.close()

    if not parsed_files:
        issues.append(
            make_issue(
                file=file_name,
                sheet="",
                row=0,
                field="",
                value=file_name,
                code="UNKNOWN_SHEET",
                message="XLSX 中没有找到任何规范工作表；工作表名必须使用 terms、periods 等 12 个名称。",
            )
        )
    return parsed_files, issues


def parse_files(files: list[tuple[str, Path]]) -> ParseResult:
    """Parse uploaded files in a deterministic order.

    ``files`` contains ``(display_name, private_path)``.  A workbook can provide
    several logical files (one per recognized worksheet); duplicate logical
    files across workbooks are rejected rather than merged.
    """
    all_files: list[ParsedFile] = []
    all_issues: list[ImportIssue] = []
    seen_logical: dict[str, str] = {}

    for display_name, path in files:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            parsed_files, issues = _read_csv(path, display_name)
        elif suffix == ".xlsx":
            parsed_files, issues = _read_xlsx(path, display_name)
        else:
            parsed_files, issues = (
                [],
                [
                    make_issue(
                        file=display_name,
                        sheet="",
                        row=0,
                        field="",
                        value=display_name,
                        code="UNSUPPORTED_FILE_TYPE",
                        message="仅支持 UTF-8 CSV 和 XLSX；请转换文件格式后重试。",
                    )
                ],
            )

        accepted_files: list[ParsedFile] = []
        for parsed_file in parsed_files:
            previous = seen_logical.get(parsed_file.logical_name)
            if previous is not None:
                all_issues.append(
                    make_issue(
                        file=display_name,
                        sheet=parsed_file.sheet,
                        row=0,
                        field="",
                        value=parsed_file.logical_name,
                        code="DUPLICATE_FILE_TYPE",
                        message=(
                            f"同一批次重复提供了 {parsed_file.logical_name}（已由 {previous} 提供）；请只保留一份。"
                        ),
                    )
                )
                continue
            seen_logical[parsed_file.logical_name] = display_name
            accepted_files.append(parsed_file)
        all_files.extend(accepted_files)
        all_issues.extend(issues)

    all_issues.sort(key=lambda item: (item.file, item.sheet, item.row, item.field, item.code))
    return ParseResult(files=all_files, issues=all_issues)


def validate_timezone(value: str) -> bool:
    try:
        ZoneInfo(value)
    except (ZoneInfoNotFoundError, ValueError):
        return False
    return True
